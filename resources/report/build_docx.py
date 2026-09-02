"""
Build the Word version of the CameraFocus documentation from the Markdown file.

Keeps the same headings, tables, code blocks and figures, so the two documents
stay in sync -- edit the Markdown, re-run this, and the .docx follows.

    python resources/report/build_docx.py

Images that are referenced but not on disk are reported and skipped, so the
document can be produced before every screenshot has been collected.
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(HERE, "CameraFocus_DOCUMENTATION.md")
DOCX = os.path.join(HERE, "CameraFocus_DOCUMENTATION.docx")
MAX_IMAGE_WIDTH = Inches(6.2)

IMAGE_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<path>[^)]+)\)\s*$")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")


def add_runs(paragraph, text):
    """Write text into a paragraph, honouring **bold** and `code` spans."""
    pos = 0
    for match in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        if match.group(1) is not None:
            paragraph.add_run(match.group(1)).bold = True
        else:
            run = paragraph.add_run(match.group(2))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_code(document, lines):
    para = document.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.space_after = Pt(8)
    run = para.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)


def add_table(document, rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = []
    for row in rows[2:]:
        cells = [c.strip() for c in row.strip("|").split("|")]
        if len(cells) < len(header):
            cells += [""] * (len(header) - len(cells))
        body.append(cells[:len(header)])
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, name in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], name)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for cells in body:
        row = table.add_row().cells
        for i, value in enumerate(cells):
            row[i].text = ""
            add_runs(row[i].paragraphs[0], value)
    document.add_paragraph()


def add_image(document, path, alt, missing):
    full = path if os.path.isabs(path) else os.path.join(HERE, path)
    if not os.path.isfile(full):
        missing.append(path)
        note = document.add_paragraph()
        run = note.add_run("[gorsel bulunamadi: %s]" % path)
        run.italic = True
        run.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    document.add_picture(full, width=MAX_IMAGE_WIDTH)
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build():
    if not os.path.isfile(MD):
        print("ERROR: markdown not found:", MD)
        return 2
    with open(MD, encoding="utf-8") as handle:
        lines = handle.read().split("\n")

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    missing = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        if line.startswith("```"):                       # code block
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(document, block)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(document, block)
            continue

        image = IMAGE_RE.match(line)
        if image:
            add_image(document, image.group("path"), image.group("alt"), missing)
            i += 1
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            if level == 1:
                document.add_heading(text, level=0)
            else:
                document.add_heading(text, level=min(level - 1, 4))
            i += 1
            continue

        if line.strip() in ("---", "***", "___"):
            document.add_page_break()
            i += 1
            continue

        stripped = line.strip()

        if stripped.startswith(("- ", "* ", "✅ ", "➖ ")):
            para = document.add_paragraph(style="List Bullet")
            add_runs(para, stripped[2:].strip())
            i += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            para = document.add_paragraph(style="List Number")
            add_runs(para, numbered.group(2))
            i += 1
            continue

        if stripped.startswith("**Şekil") or stripped.startswith("**Sekil"):
            para = document.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(para, stripped)
            for run in para.runs:
                run.italic = True
                run.font.size = Pt(9.5)
            i += 1
            continue

        para = document.add_paragraph()
        add_runs(para, stripped)
        i += 1

    document.save(DOCX)
    print("Word dokumani yazildi:", os.path.relpath(DOCX, os.path.dirname(os.path.dirname(HERE))))
    if missing:
        print("\nEksik gorseller (dosyayi bu adlarla figures/ altina koyun):")
        for path in missing:
            print("   -", path)
    return 0


if __name__ == "__main__":
    sys.exit(build())
